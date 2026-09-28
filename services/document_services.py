from __future__ import annotations

from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import (
    WD_TABLE_ALIGNMENT,
    WD_CELL_VERTICAL_ALIGNMENT,
)
from docx.shared import Inches, Pt

from fpdf import FPDF

from backend.utils.text_utils import (
    looks_like_heading,
    parse_terms,
    sanitize_text,
    split_document,
)


ROOT = Path(__file__).resolve().parents[2]

LOGO_PATH = (
    ROOT
    / "assets"
    / "legal_ease_logo.png"
)


def _add_docx_footer(
    doc: Document
) -> None:

    footer = (
        doc.sections[0]
        .footer
        .paragraphs[0]
    )

    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = footer.add_run(
        "LegalEase • AI-generated draft • Review before use"
    )

    run.font.name = "Times New Roman"
    run.font.size = Pt(8)


def format_docx(
    text: str,
    doc_type: str,
    terms: str = "",
    logo_path: str | None = None,
) -> bytes:

    doc = Document()

    section = doc.sections[0]

    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal_style = doc.styles["Normal"]

    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(11)

    chosen_logo = (
        Path(logo_path)
        if logo_path
        else LOGO_PATH
    )

    if (
        chosen_logo.exists()
        and chosen_logo.suffix.lower()
        in {".png", ".jpg", ".jpeg"}
    ):

        paragraph = doc.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        paragraph.add_run().add_picture(
            str(chosen_logo),
            width=Inches(1.2)
        )

    title = doc.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = title.add_run(
        sanitize_text(doc_type).upper()
        or "LEGAL DOCUMENT DRAFT"
    )

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for block in split_document(text):

        if looks_like_heading(block):

            paragraph = doc.add_paragraph()

            paragraph.paragraph_format.space_before = Pt(10)
            paragraph.paragraph_format.space_after = Pt(4)

            run = paragraph.add_run(block)

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

        else:

            for line in block.splitlines():

                if not line.strip():
                    continue

                paragraph = doc.add_paragraph()

                paragraph.paragraph_format.space_after = Pt(6)
                paragraph.paragraph_format.line_spacing = 1.15

                paragraph.add_run(
                    line.strip()
                )

    term_items = parse_terms(terms)

    if term_items:

        heading = doc.add_paragraph()

        run = heading.add_run(
            "KEY USER-SUPPLIED TERMS"
        )

        run.bold = True

        table = doc.add_table(
            rows=1,
            cols=2
        )

        table.alignment = (
            WD_TABLE_ALIGNMENT.CENTER
        )

        table.style = "Table Grid"

        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term"

        for index, term in enumerate(
            term_items,
            start=1
        ):

            cells = (
                table
                .add_row()
                .cells
            )

            cells[0].text = str(index)
            cells[1].text = term

            for cell in cells:
                cell.vertical_alignment = (
                    WD_CELL_VERTICAL_ALIGNMENT.CENTER
                )

    _add_docx_footer(doc)

    output = BytesIO()

    doc.save(output)

    return output.getvalue()


class BrandedPDF(FPDF):

    def __init__(
        self,
        doc_type: str
    ):

        super().__init__()

        self.doc_type = doc_type

        self.set_auto_page_break(
            auto=True,
            margin=18
        )

    def header(self):

        logo = (
            ROOT
            / "assets"
            / "legal_ease_logo.png"
        )

        if logo.exists():

            self.image(
                str(logo),
                x=80,
                y=8,
                w=50
            )

            self.ln(23)

        else:

            self.set_font(
                "Helvetica",
                "B",
                9
            )

            self.cell(
                0,
                7,
                "LegalEase",
                align="C"
            )

            self.ln(7)

    def footer(self):

        self.set_y(-14)

        self.set_font(
            "Helvetica",
            size=8
        )

        footer_text = (
            "LegalEase - "
            f"{self.doc_type[:45]} - "
            "AI-generated draft - Review before use"
        )

        self.cell(
            0,
            8,
            footer_text,
            align="C"
        )


def format_pdf(
    text: str,
    doc_type: str,
    terms: str = "",
) -> bytes:

    pdf = BrandedPDF(doc_type)

    pdf.add_page()

    pdf.set_font(
        "Helvetica",
        "B",
        16
    )

    pdf.multi_cell(
        pdf.epw,
        10,
        sanitize_text(doc_type).upper()
        or "LEGAL DOCUMENT DRAFT",
        align="C"
    )

    pdf.ln(4)

    for block in split_document(text):

        if looks_like_heading(block):

            pdf.set_font(
                "Helvetica",
                "B",
                11
            )

            pdf.multi_cell(
                pdf.epw,
                7,
                block
            )

            pdf.ln(1)

        else:

            pdf.set_font(
                "Helvetica",
                size=10.5
            )

            for line in block.splitlines():

                if line.strip():

                    pdf.multi_cell(
                        pdf.epw,
                        6,
                        line.strip()
                    )

            pdf.ln(2)

    term_items = parse_terms(terms)

    if term_items:

        pdf.set_font(
            "Helvetica",
            "B",
            11
        )

        pdf.multi_cell(
            pdf.epw,
            7,
            "KEY USER-SUPPLIED TERMS"
        )

        pdf.ln(1)

        pdf.set_font(
            "Helvetica",
            size=9.5
        )

        for index, term in enumerate(
            term_items,
            start=1
        ):

            safe_term = re.sub(
                r"[\x00-\x1f]",
                "",
                term
            )

            pdf.set_x(
                pdf.l_margin
            )

            pdf.multi_cell(
                pdf.epw,
                5.5,
                f"{index}. {safe_term}"
            )

    return bytes(
        pdf.output()
    )


def format_txt(
    text: str
) -> bytes:

    return sanitize_text(
        text
    ).encode("utf-8")